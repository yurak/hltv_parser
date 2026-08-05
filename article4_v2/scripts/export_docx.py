"""
export_docx.py — render paper/manuscript.md to paper/manuscript.docx
with journal №47 (Дніпровська політехніка) formatting:
  A4; margins 2 cm on all sides; Times New Roman;
  body 14 pt, abstracts 12 pt; line spacing 1.5; first-line indent 1.25 cm; no page numbers.
Embeds figures referenced via Markdown image syntax ![caption](../outputs/figures/x.png).

Tailored Markdown subset parser (not general): handles УДК, **meta/author lines, #/##/### headings,
paragraphs, markdown tables, image callouts, blockquote TODO notes, inline **bold** and `code`.
Math ($...$ inline, $$...$$ display) is converted to native Word equations (OMML / MS Equation)
via scripts/latex_omml.py — no more raw LaTeX text in the DOCX.
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from latex_omml import latex_to_omml

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PAPER = os.path.join(ROOT, "paper")
# Optional CLI args: [source.md] [output.docx]; default to the Ukrainian manuscript.
SRC = os.path.join(PAPER, sys.argv[1]) if len(sys.argv) > 1 else os.path.join(PAPER, "manuscript.md")
OUT = (os.path.join(PAPER, sys.argv[2]) if len(sys.argv) > 2
       else os.path.splitext(SRC)[0] + ".docx")
FONT = "Times New Roman"
BODY, ABS = 14, 12
TEXT_WIDTH = 17.0            # cm, A4 minus 2 cm margins — right tab stop for equation numbers
FIG_WIDTH = 12.5             # cm, embedded figure width (page budget: 10–20 pp.)
QUOTES = ("«", "»")          # journal typography: «ялинки» instead of straight "
CYRILLIC = re.compile(r"[а-яїієґА-ЯЇІЄҐ]")

skipped_notes = []           # internal TODO blockquotes left out of the DOCX

def smart_quotes(text):
    """Replace balanced pairs of straight quotes with the journal's «ялинки».
    Unbalanced quotes are left as-is so nothing is silently mangled."""
    return re.sub(r'"([^"\n]*)"', QUOTES[0] + r"\1" + QUOTES[1], text)

def set_lang(run, text):
    """Tag the run's proofing language so Word does not flag Ukrainian text as
    misspelled English (the document default is en-US)."""
    lang = "uk-UA" if CYRILLIC.search(text) else "en-US"
    rPr = run._element.get_or_add_rPr()
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.append(el)
    el.set(qn("w:val"), lang)

def set_font(run, size=BODY, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run._element.rPr.rFonts.set(qn("w:cs"), FONT)
    set_lang(run, run.text)

def omath_element(omml_xml):
    """Wrap a raw OMML string (from latex_omml) into an lxml element with the
    'm' namespace declared, ready to append to a paragraph's <w:p>."""
    root = "oMathPara" if omml_xml.startswith("<m:oMathPara") else "oMath"
    xml = omml_xml.replace(f"<m:{root}>", f"<m:{root} {nsdecls('m')}>", 1)
    return parse_xml(xml)

def add_runs(p, text, size=BODY):
    # first split off inline math $...$ (converted to OMML), preserving order
    for seg in re.split(r"(\$[^$]+\$)", smart_quotes(text)):
        if not seg:
            continue
        if seg.startswith("$") and seg.endswith("$") and len(seg) > 2:
            p._p.append(omath_element(latex_to_omml(seg[1:-1])))
            continue
        for tok in re.split(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)", seg):
            if not tok:
                continue
            if tok.startswith("**") and tok.endswith("**"):
                set_font(p.add_run(tok[2:-2]), size=size, bold=True)
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                set_font(p.add_run(tok[1:-1]), size=size, italic=True)
            elif tok.startswith("`") and tok.endswith("`"):
                set_font(p.add_run(tok[1:-1]), size=size)
            else:
                set_font(p.add_run(tok), size=size)

def para(doc, text, size=BODY, indent=1.25, align=None, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(indent) if indent else None
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)   # journal norm: no extra gap between paragraphs
    p.paragraph_format.space_before = Pt(0)
    # journal layout: body text justified; callers pass align explicitly for captions
    p.alignment = align if align else WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold or italic:
        set_font(p.add_run(smart_quotes(text)), size=size, bold=bold, italic=italic)
    else:
        add_runs(p, text, size=size)
    return p

def heading(doc, text, size=BODY, before=8, center=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.5
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), size=size, bold=True)
    return p

def cell_text(c, text):
    c.paragraphs[0].clear()
    t = smart_quotes(text.strip())
    bold = t.startswith("**") and t.endswith("**")
    if bold:
        t = t[2:-2]
    set_font(c.paragraphs[0].add_run(t), size=10, bold=bold)

def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            cell_text(t.cell(i, j), cell)
            if i == 0:  # header always bold
                for r in t.cell(i, j).paragraphs[0].runs:
                    r.font.bold = True

def embed_image(doc, path, caption):
    if os.path.exists(path):
        doc.add_picture(path, width=Cm(FIG_WIDTH))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, caption, size=ABS, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)

def configure_page(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2.0)

def main():
    lines = open(SRC, encoding="utf-8").read().split("\n")
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(BODY)
    configure_page(doc)

    in_abstract = False
    in_refs = False          # reference lists are set 12 pt (journal fixes 14 pt for body text only)
    eq_no = 0
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()

        if s in ("", "---"):
            i += 1; continue

        # display math block: $$ ... $$  → centered native equation
        if s == "$$":
            block = []
            i += 1
            while i < len(lines) and lines[i].strip() != "$$":
                block.append(lines[i].strip()); i += 1
            i += 1  # skip closing $$
            latex = " ".join(x for x in block if x)
            eq_no += 1
            pp = doc.add_paragraph()
            pp.paragraph_format.line_spacing = 1.5
            pp.paragraph_format.space_before = Pt(6)
            pp.paragraph_format.space_after = Pt(6)
            # centred equation + right-aligned number (N) at the right margin
            pp.paragraph_format.tab_stops.add_tab_stop(
                Cm(TEXT_WIDTH / 2), WD_TAB_ALIGNMENT.CENTER)
            pp.paragraph_format.tab_stops.add_tab_stop(
                Cm(TEXT_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
            pp.add_run().add_tab()
            pp._p.append(omath_element(latex_to_omml(latex, display=True)))
            tail = pp.add_run(f"\t({eq_no})")
            set_font(tail, size=BODY)
            continue

        # markdown image: ![caption](path)
        m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
        if m:
            cap, rel = m.group(1), m.group(2)
            embed_image(doc, os.path.normpath(os.path.join(PAPER, rel)), cap)
            i += 1; continue

        # table block
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip()); i += 1
            rows = []
            for r in block:
                if re.match(r"^\|[\s:|-]+\|$", r):
                    continue
                rows.append(r.strip("|").split("|"))
            if rows:
                add_table(doc, rows)
            continue

        # headings
        if ln.startswith("# "):
            heading(doc, ln[2:].strip(), size=BODY, before=10, center=True); i += 1; continue
        if ln.startswith("## "):
            head = ln[3:].strip()
            in_abstract = head.startswith(("Анотація", "Abstract"))
            in_refs = head.startswith(("References", "Література"))
            heading(doc, head, size=BODY, before=8); i += 1; continue
        if ln.startswith("### "):
            heading(doc, ln[4:].strip(), size=BODY, before=6); i += 1; continue

        # blockquote notes — small italic, no indent; internal TODO notes stay in the
        # Markdown source but are kept out of the submitted DOCX (reported on stdout)
        if s.startswith(">"):
            txt = re.sub(r"^>\s?", "", s)
            if txt and "TODO" in txt:
                skipped_notes.append(txt)
            elif txt:
                para(doc, txt, size=ABS, indent=0, italic=True)
            i += 1; continue

        # УДК
        if s.startswith("УДК"):
            para(doc, s, size=BODY, indent=0, bold=True); i += 1; continue

        # meta / author lines starting with **
        if s.startswith("**"):
            para(doc, s, size=(ABS if (in_abstract or "Ключові" in s or "Keywords" in s) else BODY), indent=0)
            i += 1; continue

        if in_refs:
            para(doc, s, size=ABS, indent=0.5)
            i += 1; continue

        size = ABS if in_abstract else BODY
        indent = 1.0 if in_abstract else 1.25
        para(doc, s, size=size, indent=indent)
        i += 1

    doc.save(OUT)
    print("Saved", OUT)
    if skipped_notes:
        print(f"Internal notes omitted from the DOCX ({len(skipped_notes)}):")
        for n in skipped_notes:
            print("  -", n[:160])
    leftovers = [l for l in lines if "TODO" in l and not l.strip().startswith(">")]
    if leftovers:
        print(f"WARNING: {len(leftovers)} TODO placeholder(s) still inside the manuscript body:")
        for l in leftovers:
            print("  !", l.strip()[:160])

if __name__ == "__main__":
    main()
